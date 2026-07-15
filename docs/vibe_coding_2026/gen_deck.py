# -*- coding: utf-8 -*-
"""
FocusProof — VIBE CODING 2026 — Slide thuyết trình vòng chung kết (pitching 10')
Output: VC2026-A-HT13.pptx (16 slide, 16:9)
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).parent
A = ROOT / "_assets"
OUT = ROOT / "VC2026-A-HT13.pptx"

# ---------------- Palette ----------------
INK = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x47, 0x55, 0x69)
MUTE = RGBColor(0x94, 0xA3, 0xB8)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
INDIGO_100 = RGBColor(0xE0, 0xE7, 0xFF)
VIOLET = RGBColor(0x7C, 0x3A, 0xED)
CYAN = RGBColor(0x06, 0xB6, 0xD4)
TEAL_D = RGBColor(0x0E, 0x74, 0x90)
SKY = RGBColor(0x0E, 0xA5, 0xE9)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
AMBER_D = RGBColor(0x92, 0x40, 0x0E)
GREEN = RGBColor(0x10, 0xB9, 0x81)
GREEN_D = RGBColor(0x04, 0x78, 0x57)
RED = RGBColor(0xEF, 0x44, 0x44)
PANEL = RGBColor(0xF1, 0xF5, 0xF9)
PANEL2 = RGBColor(0xF8, 0xFA, 0xFC)
BORDER = RGBColor(0xE2, 0xE8, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x0F, 0x17, 0x2A)
LIGHT_ON_DARK = RGBColor(0xC7, 0xD2, 0xFE)
CYAN_ON_DARK = RGBColor(0x7D, 0xD3, 0xFC)

FONT = "Segoe UI"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOTAL = 16


# ---------------- helpers ----------------
def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, x, y, w, h, fill, line=None, rounded=False, radius=0.08, line_w=1.0):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def oval(slide, x, y, w, h, fill, line=None, line_w=1.5):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def text(slide, x, y, w, h, s, *, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, spacing=None, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = anchor
    lines = s.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing is not None:
            p.space_after = Pt(spacing)
        r = p.add_run()
        r.text = ln
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return tb


def rich(slide, x, y, w, h, paragraphs, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=4):
    """paragraphs: list of list of (text, size, bold, color)"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = anchor
    for i, runs in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        for t, sz, bd, cl in runs:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(sz)
            r.font.bold = bd
            r.font.color.rgb = cl
    return tb


def bullets(slide, x, y, w, h, items, *, size=14, color=INK, accent=INDIGO,
            spacing=8, bullet="●", bullet_size=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        rb = p.add_run()
        rb.text = bullet + "  "
        rb.font.name = FONT
        rb.font.size = Pt(bullet_size or max(size - 4, 8))
        rb.font.color.rgb = accent
        rb.font.bold = True
        if isinstance(it, tuple):
            head, rest = it
            r1 = p.add_run()
            r1.text = head
            r1.font.name = FONT
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.color.rgb = color
            r2 = p.add_run()
            r2.text = rest
            r2.font.name = FONT
            r2.font.size = Pt(size)
            r2.font.color.rgb = SLATE
        else:
            r = p.add_run()
            r.text = it
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.color.rgb = color
    return tb


def grad_bar(slide, x, y, w, h, c1=(0x4F, 0x46, 0xE5), c2=(0x06, 0xB6, 0xD4), n=28):
    seg = int(w / n)
    for i in range(n):
        t = i / (n - 1)
        c = RGBColor(int(c1[0] + (c2[0] - c1[0]) * t),
                     int(c1[1] + (c2[1] - c1[1]) * t),
                     int(c1[2] + (c2[2] - c1[2]) * t))
        rect(slide, x + seg * i, y, seg + Emu(1), h, c)


def header(slide, num, kicker, title, subtitle=None):
    grad_bar(slide, 0, 0, SLIDE_W, Inches(0.10))
    # number chip
    rect(slide, Inches(0.6), Inches(0.38), Inches(0.62), Inches(0.62), INDIGO,
         rounded=True, radius=0.28)
    text(slide, Inches(0.6), Inches(0.38), Inches(0.62), Inches(0.62), num,
         size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, Inches(1.42), Inches(0.36), Inches(11), Inches(0.32), kicker.upper(),
         size=12, bold=True, color=TEAL_D)
    text(slide, Inches(1.40), Inches(0.62), Inches(11.3), Inches(0.62), title,
         size=27, bold=True, color=INK)
    if subtitle:
        text(slide, Inches(0.62), Inches(1.26), Inches(12.1), Inches(0.32), subtitle,
             size=12.5, color=SLATE)


def footer(slide, n):
    rect(slide, Inches(0.6), Inches(7.06), Inches(12.13), Emu(9525), BORDER)
    text(slide, Inches(0.6), Inches(7.12), Inches(9), Inches(0.3),
         "FocusProof — Chứng chỉ Tập trung Thông minh & Xác thực   ·   Đội HT13 — VIBE CODING 2026",
         size=9.5, color=MUTE)
    text(slide, Inches(11.3), Inches(7.12), Inches(1.43), Inches(0.3),
         f"{n:02d} / {TOTAL}", size=9.5, color=MUTE, align=PP_ALIGN.RIGHT)


def chip(slide, x, y, w, s, fill, color=WHITE, size=11.5, h=Inches(0.42), bold=True,
         line=None):
    rect(slide, x, y, w, h, fill, rounded=True, radius=0.5, line=line)
    text(slide, x, y, w, h, s, size=size, bold=bold, color=color,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def stat_tile(slide, x, y, w, h, number, label, accent=INDIGO, num_size=26):
    rect(slide, x, y, w, h, PANEL2, line=BORDER, rounded=True, radius=0.10)
    rect(slide, x, y, Inches(0.09), h, accent, rounded=False)
    text(slide, x + Inches(0.24), y + Inches(0.10), w - Inches(0.35), Inches(0.55),
         number, size=num_size, bold=True, color=INK)
    text(slide, x + Inches(0.24), y + h - Inches(0.52), w - Inches(0.35), Inches(0.45),
         label, size=10.5, color=SLATE)


def card(slide, x, y, w, h, title, body, accent, *, title_size=14, body_size=11.5,
         fill=PANEL2):
    rect(slide, x, y, w, h, fill, line=BORDER, rounded=True, radius=0.07)
    rect(slide, x, y, w, Inches(0.14), accent)
    text(slide, x + Inches(0.2), y + Inches(0.28), w - Inches(0.4), Inches(0.55),
         title, size=title_size, bold=True, color=INK)
    text(slide, x + Inches(0.2), y + Inches(0.82), w - Inches(0.4), h - Inches(1.0),
         body, size=body_size, color=SLATE, spacing=3)


# ======================================================================
# SLIDE 1 — BÌA
# ======================================================================
def s01_cover(prs):
    s = new_slide(prs)
    s.shapes.add_picture(str(A / "hero_bg.png"), 0, 0, SLIDE_W, SLIDE_H)

    s.shapes.add_picture(str(A / "logo_mark.png"), Inches(0.7), Inches(0.55),
                         height=Inches(0.95))
    text(s, Inches(1.85), Inches(0.72), Inches(5), Inches(0.6), "FocusProof",
         size=24, bold=True, color=WHITE)

    chip(s, Inches(0.72), Inches(2.02), Inches(5.9),
         "VIBE CODING 2026 — VÒNG CHUNG KẾT · BẢNG A · ĐỘI HT13",
         RGBColor(0x1E, 0x29, 0x3B), color=CYAN_ON_DARK, size=12, h=Inches(0.46))

    text(s, Inches(0.66), Inches(2.62), Inches(11.9), Inches(1.35), "FOCUSPROOF",
         size=66, bold=True, color=WHITE)
    text(s, Inches(0.72), Inches(3.82), Inches(11.9), Inches(0.6),
         "Biến mỗi giờ tập trung thành bằng chứng số có thể xác thực",
         size=25, bold=True, color=RGBColor(0xFB, 0xBF, 0x24))
    text(s, Inches(0.72), Inches(4.42), Inches(11.9), Inches(0.45),
         "Smart & Verifiable Focus Certificate — Chrome Extension (Manifest V3)",
         size=15, color=LIGHT_ON_DARK)

    feats = [("Chrome Extension MV3", Inches(2.30)),
             ("279 test PASS — 0 lỗi", Inches(2.30)),
             ("AI on-device · privacy-first", Inches(2.60)),
             ("Chứng chỉ ký HMAC + QR", Inches(2.55))]
    x = Inches(0.72)
    for label, w in feats:
        chip(s, x, Inches(5.30), w, label, RGBColor(0x11, 0x1C, 0x34),
             color=WHITE, size=11, h=Inches(0.44), line=RGBColor(0x33, 0x41, 0x55))
        x += w + Inches(0.22)

    rect(s, 0, Inches(6.35), SLIDE_W, Emu(18000), RGBColor(0x33, 0x41, 0x55))
    text(s, Inches(0.72), Inches(6.58), Inches(11.9), Inches(0.4),
         "Ngô Bình Minh — Đội HT13   ·   ngobinhminh2322006@gmail.com   ·   TP. Hồ Chí Minh, 07/2026",
         size=13, color=LIGHT_ON_DARK)


# ======================================================================
# SLIDE 2 — VẤN ĐỀ
# ======================================================================
def s02_problem(prs):
    s = new_slide(prs)
    header(s, "01", "Vấn đề", "Học thật — nhưng không ai tin, và không có cách chứng minh")

    # left: big stat panel (dark)
    rect(s, Inches(0.6), Inches(1.55), Inches(4.35), Inches(4.35), DARK,
         rounded=True, radius=0.05)
    text(s, Inches(0.9), Inches(1.95), Inches(3.8), Inches(1.1), "85%",
         size=64, bold=True, color=RGBColor(0xFB, 0xBF, 0x24))
    text(s, Inches(0.9), Inches(3.15), Inches(3.8), Inches(0.9),
         "sinh viên thừa nhận thường xuyên mất tập trung\nvì điện thoại & mạng xã hội",
         size=13.5, color=WHITE, spacing=3)
    rect(s, Inches(0.9), Inches(4.18), Inches(3.75), Emu(12000), RGBColor(0x33, 0x41, 0x55))
    text(s, Inches(0.9), Inches(4.38), Inches(3.8), Inches(0.6), "0",
         size=40, bold=True, color=CYAN_ON_DARK)
    text(s, Inches(0.9), Inches(5.08), Inches(3.8), Inches(0.75),
         "công cụ phổ biến nào xác thực được\ngiờ học/làm thật cho bên thứ ba",
         size=13.5, color=WHITE, spacing=3)

    # right: 3 pain cards
    pains = [
        ("Niềm tin không thể kiểm chứng", AMBER,
         "Cha mẹ, giảng viên, nhà tuyển dụng, quỹ học bổng — không ai có bằng chứng khách quan về nỗ lực học/làm việc thật."),
        ("Công cụ hiện tại chỉ đếm giờ", SKY,
         "Pomodoro, Forest, YPT… bấm giờ rồi tự khai. Treo máy đi chơi vẫn được tính là \"tập trung\" — dễ gian lận, không có dữ liệu hành vi."),
        ("Sự tập trung chưa phải tài sản", INDIGO,
         "Hàng nghìn giờ kỷ luật không đổi được thành cơ hội: học bổng, thực tập, việc làm — vì không tồn tại \"hồ sơ kỷ luật\" đáng tin."),
    ]
    y = Inches(1.55)
    for title, accent, body in pains:
        rect(s, Inches(5.25), y, Inches(7.48), Inches(1.35), PANEL2, line=BORDER,
             rounded=True, radius=0.10)
        rect(s, Inches(5.25), y, Inches(0.10), Inches(1.35), accent)
        text(s, Inches(5.52), y + Inches(0.12), Inches(7.0), Inches(0.4), title,
             size=15, bold=True, color=INK)
        text(s, Inches(5.52), y + Inches(0.53), Inches(7.05), Inches(0.75), body,
             size=11.5, color=SLATE)
        y += Inches(1.5)

    # quote strip
    rect(s, Inches(0.6), Inches(6.12), Inches(12.13), Inches(0.72), INDIGO_100,
         rounded=True, radius=0.16)
    text(s, Inches(0.95), Inches(6.12), Inches(11.6), Inches(0.72),
         "“Tôi học 8 tiếng mỗi ngày, nhưng cha mẹ vẫn hỏi: con có thật sự học không?” — trải nghiệm khởi nguồn của FocusProof",
         size=13.5, bold=True, color=INDIGO, anchor=MSO_ANCHOR.MIDDLE, italic=True)
    footer(s, 2)


# ======================================================================
# SLIDE 3 — GIẢI PHÁP
# ======================================================================
def s03_solution(prs):
    s = new_slide(prs)
    header(s, "02", "Giải pháp", "FocusProof — chứng chỉ tập trung thông minh & xác thực",
           "Chrome Extension đo tập trung bằng dữ liệu hành vi thật, rồi biến nó thành bằng chứng ai cũng kiểm chứng được")

    vals = [
        ("ĐO THẬT", INDIGO, "3 tín hiệu độc lập: khuôn mặt (AI on-device) + hoạt động chuột/bàn phím + tab đang mở."),
        ("CHẤM THẬT", SKY, "Focus Score 0–100 chấm mỗi ~6 giây theo mục tiêu phiên; cảnh báo ngay khi xao nhãng."),
        ("CHỨNG MINH ĐƯỢC", GREEN_D, "Chứng chỉ PDF + mã QR — máy chủ ký HMAC, ai quét cũng đối chiếu được bản gốc."),
        ("THẤU HIỂU", VIOLET, "AI nhận xét buổi học bằng tiếng Việt; vẫn hoạt động khi mất mạng (engine offline)."),
    ]
    y = Inches(1.78)
    for title, accent, body in vals:
        chip(s, Inches(0.6), y, Inches(2.15), title, accent, size=11.5, h=Inches(0.42))
        text(s, Inches(2.95), y - Inches(0.03), Inches(3.85), Inches(1.15), body,
             size=11.5, color=SLATE)
        y += Inches(1.24)

    # right: product mocks
    s.shapes.add_picture(str(A / "widget_mock.png"), Inches(7.05), Inches(1.72),
                         width=Inches(5.7))
    text(s, Inches(7.05), Inches(3.42), Inches(5.7), Inches(0.3),
         "Widget realtime trong lúc học — điểm tập trung cập nhật liên tục",
         size=10, color=MUTE, align=PP_ALIGN.CENTER)

    rect(s, Inches(7.05), Inches(3.92), Inches(5.7), Inches(2.75), PANEL2,
         line=BORDER, rounded=True, radius=0.06)
    text(s, Inches(7.3), Inches(4.12), Inches(5.2), Inches(0.4),
         "Một phiên học điển hình", size=14, bold=True, color=INK)
    bullets(s, Inches(7.3), Inches(4.58), Inches(5.25), Inches(2.0), [
        ("Đặt mục tiêu:  ", "“Ôn thi Giải tích 90 phút — chỉ Notion, Coursera”"),
        ("Trong phiên:  ", "widget hiển thị điểm; cảnh báo khi rời tab/mất mặt/idle"),
        ("Kết thúc:  ", "điểm + biểu đồ + AI insight + chứng chỉ PDF-QR"),
        ("Chia sẻ:  ", "gửi cha mẹ / giảng viên / đính kèm hồ sơ học bổng"),
    ], size=11.5, spacing=7, accent=INDIGO)
    footer(s, 3)


# ======================================================================
# SLIDE 4 — HÀNH TRÌNH NGƯỜI DÙNG
# ======================================================================
def s04_journey(prs):
    s = new_slide(prs)
    header(s, "03", "Sản phẩm hoạt động", "5 bước — từ mục tiêu đến bằng chứng")

    steps = [
        ("1", "Đặt mục tiêu", "Mở popup, chọn nhiệm vụ, thời lượng, domain cho phép, chế độ Strict.", INDIGO),
        ("2", "Đồng ý & bắt đầu", "Màn hình consent minh bạch (2 opt-in). Camera là tuỳ chọn — xử lý 100% tại máy.", SKY),
        ("3", "Engine chấm điểm", "3 tín hiệu hợp nhất mỗi ~6 giây. Widget nổi hiển thị điểm — cảnh báo tức thời.", CYAN),
        ("4", "Báo cáo cuối phiên", "Focus Score, biểu đồ diễn biến, số cảnh báo thật, AI insight tiếng Việt.", AMBER),
        ("5", "Chứng chỉ & chia sẻ", "PDF + QR ký HMAC phía máy chủ. Quét là ra trang xác thực công khai.", GREEN),
    ]
    for i, (n, title, desc, accent) in enumerate(steps):
        x = Inches(0.6) + Inches(2.48) * i
        oval(s, x + Inches(0.85), Inches(1.72), Inches(0.62), Inches(0.62), accent)
        text(s, x + Inches(0.85), Inches(1.72), Inches(0.62), Inches(0.62), n,
             size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        rect(s, x, Inches(2.55), Inches(2.30), Inches(2.55), PANEL2, line=BORDER,
             rounded=True, radius=0.09)
        rect(s, x, Inches(2.55), Inches(2.30), Inches(0.12), accent)
        text(s, x + Inches(0.08), Inches(2.80), Inches(2.14), Inches(0.55), title,
             size=13.5, bold=True, color=INK, align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.14), Inches(3.38), Inches(2.02), Inches(1.6), desc,
             size=10.5, color=SLATE, align=PP_ALIGN.CENTER)
        if i < 4:
            ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + Inches(2.31),
                                    Inches(1.90), Inches(0.17), Inches(0.26))
            ar.fill.solid(); ar.fill.fore_color.rgb = MUTE
            ar.line.fill.background(); ar.shadow.inherit = False

    # demo readiness strip
    rect(s, Inches(0.6), Inches(5.45), Inches(12.13), Inches(1.35), DARK,
         rounded=True, radius=0.08)
    text(s, Inches(0.95), Inches(5.62), Inches(11.5), Inches(0.42),
         "★  DEMO TRỰC TIẾP TRÊN SÂN KHẤU", size=15, bold=True,
         color=RGBColor(0xFB, 0xBF, 0x24))
    text(s, Inches(0.95), Inches(6.08), Inches(11.5), Inches(0.62),
         "Chạy phiên thật — cố tình chuyển sang Facebook để kích cảnh báo — kết thúc phiên — quét QR xác thực ngay tại chỗ.\n"
         "AI insight hoạt động cả khi đứt mạng (engine offline) · luôn có video backup phòng sự cố camera/mạng.",
         size=11.5, color=WHITE, spacing=3)
    footer(s, 4)


# ======================================================================
# SLIDE 5 — 3-SIGNAL ENGINE
# ======================================================================
def s05_engine(prs):
    s = new_slide(prs)
    header(s, "04", "Lõi công nghệ — USP", "3-Signal Focus Engine — không đối thủ nào có cùng kiến trúc")
    s.shapes.add_picture(str(A / "three_signals.png"), Inches(1.57), Inches(1.42),
                         width=Inches(10.2))
    footer(s, 5)


# ======================================================================
# SLIDE 6 — KIẾN TRÚC
# ======================================================================
def s06_architecture(prs):
    s = new_slide(prs)
    header(s, "05", "Kiến trúc hệ thống",
           "Xử lý tại thiết bị — xác thực tại máy chủ",
           "React 18 · TypeScript · Vite · MediaPipe BlazeFace (WASM) · Node/Express · SQLite · GitHub Actions CI")
    s.shapes.add_picture(str(A / "architecture.png"), Inches(1.72), Inches(1.60),
                         width=Inches(9.9))
    footer(s, 6)


# ======================================================================
# SLIDE 7 — CHỨNG CHỈ XÁC THỰC (CHỐNG GIAN LẬN)
# ======================================================================
def s07_certificate(prs):
    s = new_slide(prs)
    header(s, "06", "Điểm khác biệt cốt lõi", "Chứng chỉ không thể giả mạo — vì máy chủ ký từng bản ghi")

    flow = [
        ("Hash tại máy", "Kết thúc phiên, dữ liệu điểm được băm SHA-256 — \"dấu vân tay\" toàn vẹn.", INDIGO),
        ("Máy chủ ký HMAC", "Bản ghi (điểm, hạng, hash) được ký HMAC-SHA-256 v2 bằng khoá bí mật server — chống ghi đè (409).", VIOLET),
        ("PDF + QR", "Chứng chỉ nhúng QR trỏ tới trang xác thực công khai /verify/:id.", SKY),
        ("Ai cũng kiểm chứng được", "Cha mẹ / giảng viên / nhà tuyển dụng quét QR → đối chiếu bản gốc do server lưu. Sửa 1 ký tự là hash mất hiệu lực.", GREEN_D),
    ]
    y = Inches(1.62)
    for i, (title, body, accent) in enumerate(flow):
        oval(s, Inches(0.68), y + Inches(0.10), Inches(0.44), Inches(0.44), accent)
        text(s, Inches(0.68), y + Inches(0.10), Inches(0.44), Inches(0.44), str(i + 1),
             size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(1.30), y, Inches(4.9), Inches(0.38), title, size=14.5,
             bold=True, color=INK)
        text(s, Inches(1.30), y + Inches(0.37), Inches(4.9), Inches(0.85), body,
             size=11, color=SLATE)
        if i < 3:
            rect(s, Inches(0.88), y + Inches(0.62), Emu(28000), Inches(0.60), BORDER)
        y += Inches(1.28)

    rect(s, Inches(0.6), Inches(6.30), Inches(5.75), Inches(0.55), INDIGO_100,
         rounded=True, radius=0.2)
    text(s, Inches(0.82), Inches(6.30), Inches(5.4), Inches(0.55),
         "Trung thực kỹ thuật: SHA-256 = toàn vẹn dữ liệu · chữ ký thật = HMAC phía máy chủ",
         size=10.5, bold=True, color=INDIGO, anchor=MSO_ANCHOR.MIDDLE)

    s.shapes.add_picture(str(A / "certificate_mock.png"), Inches(6.62), Inches(1.62),
                         width=Inches(6.15))
    text(s, Inches(6.62), Inches(5.95), Inches(6.15), Inches(0.3),
         "Chứng chỉ minh hoạ — QR thật, quét được", size=10, color=MUTE,
         align=PP_ALIGN.CENTER)
    footer(s, 7)


# ======================================================================
# SLIDE 8 — BẢO MẬT & QUYỀN RIÊNG TƯ
# ======================================================================
def s08_privacy(prs):
    s = new_slide(prs)
    header(s, "07", "Bảo mật & Quyền riêng tư", "Privacy-first không phải khẩu hiệu — là kiến trúc")

    # left dark: 5 KHÔNG
    rect(s, Inches(0.6), Inches(1.6), Inches(5.6), Inches(5.2), DARK,
         rounded=True, radius=0.05)
    text(s, Inches(0.95), Inches(1.85), Inches(5.0), Inches(0.5),
         "5 ĐIỀU CHÚNG TÔI KHÔNG LÀM", size=17, bold=True, color=RGBColor(0xFB, 0xBF, 0x24))
    nos = [
        "KHÔNG ghi hình — không một frame ảnh nào rời thiết bị",
        "KHÔNG đọc cảm xúc / trạng thái tâm lý (tuân thủ tinh thần EU AI Act)",
        "KHÔNG thu nội dung gõ phím mặc định; không bao giờ đọc ô mật khẩu / OTP / thẻ",
        "KHÔNG nhúng API key vào client — key chỉ nằm ở server (CI chặn secret lộ)",
        "KHÔNG theo dõi âm thầm — người học tự bấm bắt đầu, thấy rõ widget suốt phiên",
    ]
    yy = Inches(2.42)
    for t in nos:
        oval(s, Inches(0.95), yy + Inches(0.03), Inches(0.30), Inches(0.30), RED)
        text(s, Inches(0.95), yy - Inches(0.015), Inches(0.30), Inches(0.36), "✕",
             size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(1.42), yy - Inches(0.05), Inches(4.6), Inches(0.85), t,
             size=11.5, color=WHITE)
        yy += Inches(0.86)

    # right: measures
    cards_ = [
        ("Consent & minh bạch", GREEN_D,
         "Màn hình đồng ý trước phiên đầu (chuẩn Chrome Web Store): liệt kê đúng dữ liệu thu thập, camera & nội dung gõ là 2 opt-in riêng."),
        ("AI chạy tại thiết bị", INDIGO,
         "BlazeFace WASM nhận diện có/không khuôn mặt ngay trên trình duyệt. Chỉ điểm tổng hợp được lưu — không hình ảnh."),
        ("Phòng thủ nhiều lớp", SKY,
         "CSP + DOMPurify · domain matching theo hostname (chống fake-notion.so.evil.com) · token API · giải trình 7 quyền trong tài liệu nộp kèm."),
    ]
    y = Inches(1.6)
    for title, accent, body in cards_:
        card(s, Inches(6.45), y, Inches(6.28), Inches(1.62), title, body, accent,
             body_size=11)
        y += Inches(1.79)
    footer(s, 8)


# ======================================================================
# SLIDE 9 — CHẤT LƯỢNG KỸ THUẬT
# ======================================================================
def s09_quality(prs):
    s = new_slide(prs)
    header(s, "08", "Chất lượng kỹ thuật", "Mọi con số đều kiểm chứng lại được từ gói nộp")
    s.shapes.add_picture(str(A / "quality_dashboard.png"), Inches(0.62), Inches(1.55),
                         width=Inches(12.1))
    rect(s, Inches(0.6), Inches(6.25), Inches(12.13), Inches(0.62), INDIGO_100,
         rounded=True, radius=0.16)
    text(s, Inches(0.95), Inches(6.25), Inches(11.6), Inches(0.62),
         "10/10 góp ý của giám khảo vòng loại đã xử lý — mỗi lỗi sửa xong đều kèm test hồi quy riêng (+95 test mới)",
         size=13, bold=True, color=INDIGO, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 9)


# ======================================================================
# SLIDE 10 — QUY TRÌNH VIBE CODING
# ======================================================================
def s10_process(prs):
    s = new_slide(prs)
    header(s, "09", "Quy trình phát triển", "Vibe coding có kỷ luật — AI viết, con người kiểm chứng")

    steps = [
        ("Spec & kiến trúc", "Con người quyết định bài toán, ràng buộc, luồng dữ liệu, tiêu chí nghiệm thu.", INK),
        ("AI pair-coding", "Claude Code sinh mã theo spec — tốc độ của cả một đội dev.", INDIGO),
        ("Test hồi quy", "Vitest — 279 test; mỗi bug từng gặp có test riêng chống tái diễn.", SKY),
        ("Audit đa tác tử", "30 AI agent độc lập đọc toàn bộ mã, chấm theo rubric, săn lỗi bảo mật/logic.", VIOLET),
        ("Cổng CI/CD", "GitHub Actions + husky: lint, type-check, test, build, chặn secret — trước mỗi commit.", GREEN_D),
    ]
    y = Inches(1.62)
    for i, (title, desc, accent) in enumerate(steps):
        rect(s, Inches(0.6), y, Inches(6.4), Inches(0.94), PANEL2, line=BORDER,
             rounded=True, radius=0.12)
        rect(s, Inches(0.6), y, Inches(0.09), Inches(0.94), accent)
        oval(s, Inches(0.82), y + Inches(0.25), Inches(0.44), Inches(0.44), accent)
        text(s, Inches(0.82), y + Inches(0.25), Inches(0.44), Inches(0.44), str(i + 1),
             size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(1.45), y + Inches(0.08), Inches(5.4), Inches(0.36), title,
             size=13.5, bold=True, color=INK)
        text(s, Inches(1.45), y + Inches(0.44), Inches(5.45), Inches(0.48), desc,
             size=10.5, color=SLATE)
        y += Inches(1.06)

    s.shapes.add_picture(str(A / "radar_rubric.png"), Inches(7.42), Inches(1.52),
                         height=Inches(5.35))
    footer(s, 10)


# ======================================================================
# SLIDE 11 — CẠNH TRANH
# ======================================================================
def s11_compare(prs):
    s = new_slide(prs)
    header(s, "10", "Vị thế cạnh tranh", "Đứng ở giao điểm mà chưa sản phẩm nào chạm tới")

    cols = [Inches(4.00), Inches(1.95), Inches(1.95), Inches(1.95), Inches(2.28)]
    heads = ["Tiêu chí", "Forest / Pomodoro", "YPT (5tr+ user)", "Azota", "FOCUSPROOF"]
    rows = [
        ("Đo tập trung thật (hành vi)", "✕  chỉ đếm giờ", "✕  tự khai", "✕  chỉ lúc thi", "✓  3 tín hiệu"),
        ("Bằng chứng xác thực được", "✕", "✕", "△  biên bản thi", "✓  QR + HMAC"),
        ("AI insight tiếng Việt", "✕", "✕", "✕", "✓  offline-safe"),
        ("Riêng tư: AI on-device", "—  không camera", "—  không camera", "✕  giám sát", "✓  0 ảnh rời máy"),
        ("Người học sở hữu dữ liệu", "✕", "✓", "✕  trường quản lý", "✓  opt-in toàn phần"),
    ]

    x0, y0 = Inches(0.6), Inches(1.58)
    rh, hh = Inches(0.78), Inches(0.60)
    # header row
    x = x0
    for j, htxt in enumerate(heads):
        fill = INDIGO if j == 4 else DARK
        rect(s, x, y0, cols[j] - Emu(22000), hh, fill, rounded=False)
        text(s, x, y0, cols[j] - Emu(22000), hh, htxt, size=11.5, bold=True,
             color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += cols[j]
    # body
    for i, row in enumerate(rows):
        y = y0 + hh + rh * i + Emu(18000) * (i + 1)
        x = x0
        for j, val in enumerate(row):
            if j == 0:
                rect(s, x, y, cols[j] - Emu(22000), rh, PANEL, rounded=False)
                text(s, x + Inches(0.14), y, cols[j] - Inches(0.2), rh, val,
                     size=11.5, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
            else:
                fill = INDIGO_100 if j == 4 else PANEL2
                rect(s, x, y, cols[j] - Emu(22000), rh, fill, line=BORDER)
                good = val.startswith("✓")
                bad = val.startswith("✕")
                c = GREEN_D if good else (SLATE if not bad else RGBColor(0x9F, 0x1D, 0x1D))
                text(s, x, y, cols[j] - Emu(22000), rh, val, size=10.5,
                     bold=good, color=c, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            x += cols[j]

    rect(s, Inches(0.6), Inches(6.22), Inches(12.13), Inches(0.66), DARK,
         rounded=True, radius=0.16)
    text(s, Inches(0.95), Inches(6.22), Inches(11.6), Inches(0.66),
         "“Azota chứng minh bạn không gian lận lúc thi — FocusProof chứng minh bạn có nỗ lực lúc học.”",
         size=14, bold=True, color=CYAN_ON_DARK, anchor=MSO_ANCHOR.MIDDLE,
         align=PP_ALIGN.CENTER, italic=True)
    footer(s, 11)


# ======================================================================
# SLIDE 12 — THỊ TRƯỜNG
# ======================================================================
def s12_market(prs):
    s = new_slide(prs)
    header(s, "11", "Thị trường", "Bắt đầu từ 2,3 triệu sinh viên Việt Nam",
           "Sản phẩm born-global: giao diện song ngữ, phân phối qua Chrome Web Store toàn cầu ngay từ MVP")
    s.shapes.add_picture(str(A / "market_chart.png"), Inches(1.07), Inches(1.78),
                         width=Inches(11.2))
    footer(s, 12)


# ======================================================================
# SLIDE 13 — MÔ HÌNH KINH DOANH
# ======================================================================
def s13_business(prs):
    s = new_slide(prs)
    header(s, "12", "Mô hình kinh doanh", "Freemium + Credit — không phụ thuộc quảng cáo")

    plans = [
        ("FREE", "0 đ", "Đo tập trung + chứng chỉ cơ bản.\nCửa vào không rào cản — nguồn viral.", SLATE),
        ("PRO", "49K /tháng", "AI insight không giới hạn, lịch sử\nđầy đủ, chứng chỉ ký server.", INDIGO),
        ("TEAM", "199K /tháng·5ng", "Nhóm học tập / freelancer:\nbáo cáo chung, leaderboard riêng.", VIOLET),
        ("TRƯỜNG HỌC", "499K+ /tháng", "Dashboard giảng viên, join-code lớp,\nemail report tự động.", TEAL_D),
    ]
    for i, (name, price, desc, accent) in enumerate(plans):
        x = Inches(0.6) + Inches(3.10) * i
        rect(s, x, Inches(1.62), Inches(2.90), Inches(2.42), PANEL2, line=BORDER,
             rounded=True, radius=0.08)
        rect(s, x, Inches(1.62), Inches(2.90), Inches(0.5), accent)
        text(s, x, Inches(1.62), Inches(2.90), Inches(0.5), name, size=13.5,
             bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x, Inches(2.26), Inches(2.90), Inches(0.55), price, size=21,
             bold=True, color=INK, align=PP_ALIGN.CENTER)
        text(s, x + Inches(0.15), Inches(2.90), Inches(2.60), Inches(1.05), desc,
             size=10.5, color=SLATE, align=PP_ALIGN.CENTER, spacing=3)
    text(s, Inches(0.6), Inches(4.18), Inches(12.13), Inches(0.35),
         "+ Gói credit 19K–99K cho người dùng thỉnh thoảng cần AI/PDF — server đếm credit, chống lạm dụng.",
         size=11, color=SLATE, align=PP_ALIGN.CENTER)

    tiles = [
        ("CAC ≤ 30K", "mỗi chứng chỉ chia sẻ là một\nkênh marketing tự nhiên", AMBER),
        ("LTV ≈ 588K", "12 tháng × 49K —\ngiữ chân bằng streak & hồ sơ", SKY),
        ("LTV/CAC ≈ 19×", "đơn vị kinh tế lành mạnh\nngay từ mô hình cơ sở", INDIGO),
        ("Chi phí biên ≈ 0", "AI chạy tại thiết bị người dùng ·\nSQLite/free-tier cho backend", GREEN_D),
    ]
    for i, (num, label, accent) in enumerate(tiles):
        x = Inches(0.6) + Inches(3.10) * i
        stat_tile(s, x, Inches(4.72), Inches(2.90), Inches(1.30), num, label,
                  accent=accent, num_size=22)

    text(s, Inches(0.6), Inches(6.30), Inches(12.13), Inches(0.5),
         "Năm 1 (mục tiêu): 10.000 lượt cài · 1.500 người trả phí · ≈ 880 triệu đồng doanh thu · hoà vốn tháng 8–9",
         size=13, bold=True, color=INK, align=PP_ALIGN.CENTER)
    footer(s, 13)


# ======================================================================
# SLIDE 14 — LỘ TRÌNH
# ======================================================================
def s14_roadmap(prs):
    s = new_slide(prs)
    header(s, "13", "Lộ trình phát triển", "Sản phẩm đã chạy được hôm nay — không phải lời hứa")

    milestones = [
        ("HÔM NAY", "MVP 1.0.1 hoàn chỉnh", "279 test PASS · gói nộp 7.7MB ·\nbackend SQLite + trang verify ·\nsẵn sàng demo trực tiếp", GREEN_D, True),
        ("Q3 / 2026", "Ra mắt công khai", "Chrome Web Store (unlisted → public) ·\ndeploy backend production ·\nđo CPU/RAM phiên dài công bố kèm", INDIGO, False),
        ("Q4 / 2026", "Chiều lớp học", "Dashboard giảng viên · join-code lớp ·\nemail report tự động sau phiên ·\nthí điểm 3 lớp tại TDTU", VIOLET, False),
        ("2027", "Hệ sinh thái", "Google Classroom / LMS ·\nhồ sơ kỷ luật cá nhân theo năm ·\nmở rộng Đông Nam Á", SKY, False),
    ]
    # timeline line
    rect(s, Inches(1.0), Inches(2.42), Inches(11.3), Emu(28000), BORDER)
    for i, (when, title, desc, accent, done) in enumerate(milestones):
        x = Inches(0.75) + Inches(3.06) * i
        cx = x + Inches(1.30)
        oval(s, cx, Inches(2.27), Inches(0.36), Inches(0.36), accent)
        if done:
            text(s, cx, Inches(2.25), Inches(0.36), Inches(0.40), "✓", size=13,
                 bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        chip(s, x + Inches(0.53), Inches(1.62), Inches(1.90), when, accent,
             size=12, h=Inches(0.44))
        rect(s, x, Inches(3.02), Inches(2.96), Inches(2.55), PANEL2, line=BORDER,
             rounded=True, radius=0.08)
        text(s, x + Inches(0.18), Inches(3.20), Inches(2.62), Inches(0.65), title,
             size=14.5, bold=True, color=INK)
        text(s, x + Inches(0.18), Inches(3.80), Inches(2.66), Inches(1.7), desc,
             size=10.5, color=SLATE, spacing=3)

    rect(s, Inches(0.6), Inches(6.02), Inches(12.13), Inches(0.80), INDIGO_100,
         rounded=True, radius=0.12)
    text(s, Inches(0.95), Inches(6.02), Inches(11.6), Inches(0.80),
         "Định vị nhất quán: công cụ TỰ CHỨNG MINH do người học sở hữu (opt-in) — tuyệt đối không đi vào vùng proctoring/giám sát áp đặt,\nvùng đã bị phản ứng dữ dội trên thế giới và bị EU AI Act giới hạn.",
         size=11.5, bold=True, color=INDIGO, anchor=MSO_ANCHOR.MIDDLE, spacing=3)
    footer(s, 14)


# ======================================================================
# SLIDE 15 — ĐỘI NGŨ
# ======================================================================
def s15_team(prs):
    s = new_slide(prs)
    header(s, "14", "Đội ngũ", "Đội HT13 — nhỏ về số người, đủ về năng lực")

    # left card
    rect(s, Inches(0.6), Inches(1.62), Inches(3.7), Inches(4.95), PANEL2,
         line=BORDER, rounded=True, radius=0.06)
    rect(s, Inches(0.6), Inches(1.62), Inches(3.7), Inches(0.5), DARK)
    text(s, Inches(0.6), Inches(1.62), Inches(3.7), Inches(0.5), "ĐỘI HT13 — BẢNG A",
         size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    oval(s, Inches(1.55), Inches(2.45), Inches(1.8), Inches(1.8), INDIGO, line=CYAN,
         line_w=2.5)
    text(s, Inches(1.55), Inches(2.45), Inches(1.8), Inches(1.8), "NBM", size=40,
         bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(0.7), Inches(4.42), Inches(3.5), Inches(0.42), "Ngô Bình Minh",
         size=18, bold=True, color=INK, align=PP_ALIGN.CENTER)
    text(s, Inches(0.7), Inches(4.85), Inches(3.5), Inches(0.7),
         "MSSV 524H0169 — Khoa CNTT\nĐH Tôn Đức Thắng", size=12, color=SLATE,
         align=PP_ALIGN.CENTER, spacing=3)
    text(s, Inches(0.7), Inches(5.62), Inches(3.5), Inches(0.75),
         "Founder · Full-stack · AI Integrator", size=12.5, bold=True, color=TEAL_D,
         align=PP_ALIGN.CENTER)

    # right: roles
    roles = [
        ("Product Owner", "định hình bài toán, rubric nghiệm thu, lộ trình", INDIGO),
        ("Full-stack Developer", "Chrome Extension MV3, React, Node/Express, SQLite", SKY),
        ("AI Engineer", "BlazeFace on-device, GPT-4o-mini proxy, engine offline", VIOLET),
        ("QA & DevOps", "279 test, CI/CD, audit đa tác tử, check-no-secrets", GREEN_D),
        ("Design & Growth", "UI/UX popup-widget-chứng chỉ, viral loop chia sẻ", AMBER_D),
    ]
    y = Inches(1.62)
    for title, desc, accent in roles:
        rect(s, Inches(4.6), y, Inches(8.13), Inches(0.80), PANEL2, line=BORDER,
             rounded=True, radius=0.14)
        rect(s, Inches(4.6), y, Inches(0.09), Inches(0.80), accent)
        text(s, Inches(4.88), y + Inches(0.075), Inches(2.9), Inches(0.65), title,
             size=13, bold=True, color=INK, anchor=MSO_ANCHOR.MIDDLE)
        text(s, Inches(7.85), y + Inches(0.075), Inches(4.75), Inches(0.65), desc,
             size=11, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(0.945)

    rect(s, Inches(4.6), Inches(6.30), Inches(8.13), Inches(0.55), DARK,
         rounded=True, radius=0.2)
    text(s, Inches(4.85), Inches(6.30), Inches(7.7), Inches(0.55),
         "1 người + AI = năng suất của một đội — nhưng kiểm chứng nghiêm ngặt như một phòng QA.",
         size=11.5, bold=True, color=CYAN_ON_DARK, anchor=MSO_ANCHOR.MIDDLE)
    footer(s, 15)


# ======================================================================
# SLIDE 16 — KẾT
# ======================================================================
def s16_closing(prs):
    s = new_slide(prs)
    s.shapes.add_picture(str(A / "hero_bg.png"), 0, 0, SLIDE_W, SLIDE_H)

    s.shapes.add_picture(str(A / "logo_mark.png"), Inches(0.72), Inches(0.85),
                         height=Inches(1.05))
    text(s, Inches(0.72), Inches(2.25), Inches(11.9), Inches(1.0), "FOCUSPROOF",
         size=52, bold=True, color=WHITE)
    text(s, Inches(0.74), Inches(3.25), Inches(11.6), Inches(0.95),
         "Tập trung không còn là cảm giác —\nmà là một bằng chứng có thể xác thực.",
         size=27, bold=True, color=RGBColor(0xFB, 0xBF, 0x24), spacing=4)

    chips_ = [("Demo trực tiếp ngay sau đây", Inches(2.85)),
              ("Prototype hoàn chỉnh — 279 test PASS", Inches(3.45)),
              ("Sẵn sàng Chrome Web Store", Inches(2.90))]
    x = Inches(0.74)
    for label, w in chips_:
        chip(s, x, Inches(4.72), w, label, RGBColor(0x11, 0x1C, 0x34), color=WHITE,
             size=11.5, h=Inches(0.46), line=RGBColor(0x33, 0x41, 0x55))
        x += w + Inches(0.25)

    text(s, Inches(0.74), Inches(5.55), Inches(11.9), Inches(0.4),
         "Kính mời Ban giám khảo trải nghiệm một phiên tập trung — và tự tay quét QR xác thực chứng chỉ.",
         size=14.5, color=LIGHT_ON_DARK)

    rect(s, 0, Inches(6.35), SLIDE_W, Emu(18000), RGBColor(0x33, 0x41, 0x55))
    text(s, Inches(0.74), Inches(6.58), Inches(11.9), Inches(0.4),
         "Đội HT13 — Ngô Bình Minh   ·   ngobinhminh2322006@gmail.com   ·   VIBE CODING 2026 — Vòng chung kết",
         size=13, color=LIGHT_ON_DARK)


# ======================================================================
def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    builders = [s01_cover, s02_problem, s03_solution, s04_journey, s05_engine,
                s06_architecture, s07_certificate, s08_privacy, s09_quality,
                s10_process, s11_compare, s12_market, s13_business, s14_roadmap,
                s15_team, s16_closing]
    for b in builders:
        b(prs)

    core = prs.core_properties
    core.title = "FocusProof — VIBE CODING 2026 — Đội HT13"
    core.author = "Đội HT13 — Ngô Bình Minh"
    core.subject = "Slide thuyết trình vòng chung kết"

    prs.save(OUT)
    print(f"[OK] {OUT.name} — {len(prs.slides)} slides, {OUT.stat().st_size // 1024} KB")


if __name__ == "__main__":
    main()
